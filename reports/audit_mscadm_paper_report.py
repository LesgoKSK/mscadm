from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, is_zipfile

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports" / "MS-CADM_论文方法与复现专题报告.docx"
SOURCE = ROOT / "reports" / "MSCADM_PAPER_REPRODUCTION_REPORT.md"
OUTPUT = ROOT / "reports" / "qa" / "mscadm_paper_audit.json"
TABLE1 = ROOT / "outputs" / "full_reproduction" / "tables" / "table1_metrics.csv"
PAPER_CONFIG = ROOT / "repro_configs" / "paper.json"
ASSETS = [
    ROOT / "reports" / "assets" / "mscadm_paper" / "diffusion_workflow.png",
    ROOT / "reports" / "assets" / "mscadm_paper" / "mscadm_architecture.png",
    ROOT / "reports" / "assets" / "mscadm_paper" / "training_curve.png",
    ROOT / "reports" / "assets" / "mscadm_paper" / "reproduction_scorecard.png",
    ROOT / "outputs" / "full_reproduction" / "figures" / "figure4_mscadm_test_250steps.png",
    ROOT / "outputs" / "full_reproduction" / "figures" / "figure5_intervals.png",
    ROOT / "outputs" / "full_reproduction" / "figures" / "figure7_distribution.png",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def check(condition: bool, name: str, details: object, checks: list[dict]) -> None:
    checks.append({"name": name, "passed": bool(condition), "details": details})


def main() -> None:
    checks: list[dict] = []
    check(REPORT.exists() and REPORT.stat().st_size > 100_000, "report_exists", REPORT.stat().st_size, checks)
    check(is_zipfile(REPORT), "valid_openxml_zip", str(REPORT), checks)

    document = Document(REPORT)
    paragraphs = [p.text for p in document.paragraphs]
    table_text = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    full_text = "\n".join(paragraphs + table_text)

    check(len(document.paragraphs) == 290, "paragraph_count", len(document.paragraphs), checks)
    check(len(document.tables) == 25, "table_count", len(document.tables), checks)
    check(len(document.inline_shapes) == 7, "figure_count", len(document.inline_shapes), checks)
    check(len(document.sections) >= 2, "cover_and_body_sections", len(document.sections), checks)

    headings = {
        style: [p.text for p in document.paragraphs if p.style and p.style.name == style]
        for style in ("Heading 1", "Heading 2", "Heading 3")
    }
    check(len(headings["Heading 1"]) >= 10, "heading_1_count", headings["Heading 1"], checks)
    required_h1 = {
        "1. 论文任务、贡献与原始主张",
        "4. MS-CADM 网络架构与张量流",
        "8. 复现结果",
        "9. 未复现原因和原文可复现性漏洞",
        "11. 最终评价",
    }
    check(required_h1.issubset(set(headings["Heading 1"])), "required_sections", sorted(required_h1), checks)
    suc_heading = next((p for p in document.paragraphs if p.text.startswith("8.7 Table 5：RTS-24 SUC 重构")), None)
    check(suc_heading is not None and suc_heading.paragraph_format.page_break_before is True,
          "suc_subsection_starts_new_page", suc_heading.text if suc_heading else None, checks)

    forbidden = ("[[COVER_END]]", "[[END_REPORT]]", "[[FIGURE:", "[[CUSTOM_FIGURE:", "�")
    found_forbidden = [token for token in forbidden if token in full_text]
    check(not found_forbidden, "no_unresolved_markers_or_replacement_chars", found_forbidden, checks)

    critical_strings = [
        "1,478,018",
        "0.124248",
        "0.177632",
        "0.107400",
        "0.054139",
        "0.681950",
        "22.0200",
        "0.3595",
        "26,000",
        "395,386.19",
        "pooled zone-day",
        "Algorithm 2",
        "learned variance",
    ]
    missing_critical = [value for value in critical_strings if value not in full_text]
    check(not missing_critical, "critical_content_present", missing_critical, checks)

    config = json.loads(PAPER_CONFIG.read_text(encoding="utf-8"))["models"]["mscadm"]
    locked = {
        "model_dim": config["model"]["model_dim"],
        "condition_bottleneck": config["model"]["condition_bottleneck"],
        "depth": config["model"]["depth"],
        "heads": config["model"]["heads"],
        "ff_multiplier": config["model"]["ff_multiplier"],
        "timesteps": config["diffusion"]["timesteps"],
        "steps": config["training"]["steps"],
        "batch_size": config["training"]["batch_size"],
    }
    config_strings = ["128", "64", "250", "26,000", "256"]
    check(all(value in full_text for value in config_strings), "config_values_present", locked, checks)

    with TABLE1.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    headline = next(row for row in rows if row["model"] == "mscadm_test_250steps")
    formatted = {
        "MAE": f"{float(headline['MAE']):.6f}",
        "RMSE": f"{float(headline['RMSE']):.6f}",
        "CRPS": f"{float(headline['CRPS']):.6f}",
        "QS": f"{float(headline['QS']):.6f}",
        "ES": f"{float(headline['ES']):.6f}",
        "VS": f"{float(headline['VS']):.4f}",
    }
    check(all(value in full_text for value in formatted.values()), "headline_matches_table1_archive", formatted, checks)

    ragged = []
    for index, table in enumerate(document.tables, start=1):
        counts = [len(row.cells) for row in table.rows]
        if len(set(counts)) != 1:
            ragged.append({"table": index, "counts": counts})
    check(not ragged, "no_ragged_tables", ragged, checks)

    with ZipFile(REPORT) as archive:
        names = archive.namelist()
        document_xml = archive.read("word/document.xml")
        media = sorted(name for name in names if name.startswith("word/media/") and not name.endswith("/"))
    check(len(media) == 7, "embedded_media_count", media, checks)

    root = document.element
    ns = root.nsmap
    doc_prs = root.xpath(".//wp:docPr")
    alt_records = [
        {"title": node.get("title", ""), "description": node.get("descr", "")}
        for node in doc_prs
    ]
    check(len(alt_records) == 7 and all(x["title"] and x["description"] for x in alt_records),
          "all_figures_have_alt_text", alt_records, checks)

    geometry = []
    geometry_ok = True
    for index, table in enumerate(root.xpath(".//w:tbl"), start=1):
        grid = table.xpath("./w:tblGrid/w:gridCol")
        widths = [int(node.get(qn("w:w"))) for node in grid]
        tbl_w = table.xpath("./w:tblPr/w:tblW")
        table_width = int(tbl_w[0].get(qn("w:w"))) if tbl_w else None
        indent = table.xpath("./w:tblPr/w:tblInd")
        table_indent = int(indent[0].get(qn("w:w"))) if indent else None
        header = bool(table.xpath("./w:tr[1]/w:trPr/w:tblHeader"))
        cant_split_rows = len(table.xpath("./w:tr/w:trPr/w:cantSplit"))
        row_count = len(table.xpath("./w:tr"))
        ok = sum(widths) == 9360 and table_width == 9360 and table_indent == 120 and header and cant_split_rows == row_count
        geometry_ok = geometry_ok and ok
        geometry.append({
            "table": index,
            "grid_sum": sum(widths),
            "table_width": table_width,
            "indent": table_indent,
            "header_repeats": header,
            "cant_split_rows": f"{cant_split_rows}/{row_count}",
        })
    check(geometry_ok, "table_geometry_and_pagination", geometry, checks)

    section_details = []
    letter_ok = True
    for section in document.sections:
        item = {
            "page_width": section.page_width,
            "page_height": section.page_height,
            "top_margin": section.top_margin,
            "bottom_margin": section.bottom_margin,
            "left_margin": section.left_margin,
            "right_margin": section.right_margin,
        }
        section_details.append(item)
        letter_ok = letter_ok and abs(section.page_width.inches - 8.5) < 0.01 and abs(section.page_height.inches - 11.0) < 0.01
    check(letter_ok, "letter_page_geometry", section_details, checks)

    missing_assets = [str(path) for path in ASSETS if not path.exists() or path.stat().st_size == 0]
    check(not missing_assets, "source_assets_present", missing_assets, checks)
    check(document.core_properties.title == "MS-CADM 论文方法与复现专题报告",
          "core_title", document.core_properties.title, checks)

    passed = all(item["passed"] for item in checks)
    result = {
        "passed": passed,
        "report": str(REPORT),
        "sha256": sha256(REPORT),
        "source_sha256": sha256(SOURCE),
        "size_bytes": REPORT.stat().st_size,
        "checks": checks,
        "notes": [
            "LibreOffice/soffice is unavailable, so the official render_docx.py conversion could not run.",
            "A 30-page Chrome/HTML surrogate preview was rasterized and inspected separately; this is layout QA, not a claim of byte-identical Word rendering.",
        ],
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"passed": passed, "checks": len(checks), "sha256": result["sha256"], "output": str(OUTPUT)}, ensure_ascii=False))
    if not passed:
        failed = [item for item in checks if not item["passed"]]
        print(json.dumps(failed, ensure_ascii=False, indent=2))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
