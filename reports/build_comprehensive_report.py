from __future__ import annotations

import re
import os
import sys
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "COMPREHENSIVE_RESEARCH_REPORT.md"
OUTPUT = ROOT / "reports" / "MS-CADM_复现与改进实验综合研究报告.docx"
ASSET_DIR = ROOT / "reports" / "assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

DOC_SKILL_VALUE = os.environ.get("DOCUMENT_SKILL_DIR")
if not DOC_SKILL_VALUE:
    raise RuntimeError(
        "Set DOCUMENT_SKILL_DIR to the document-skill directory that contains scripts/table_geometry.py"
    )
DOC_SKILL = Path(DOC_SKILL_VALUE)
sys.path.insert(0, str(DOC_SKILL / "scripts"))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GOLD = "C59A3D"
TEXT = "263238"
MUTED = "66737F"
LIGHT = "F4F6F9"
LIGHT_BLUE = "EAF2F8"
GREEN = "2E7D32"
AMBER = "C27C0E"
RED = "B23A48"
BLUE_GREY = "607D8B"
WHITE = "FFFFFF"


def set_run_font(run, *, ascii_name="Calibri", east_asia="Microsoft YaHei", size=None,
                 color=None, bold=None, italic=None):
    run.font.name = ascii_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, *, ascii_name="Calibri", east_asia="Microsoft YaHei", size=None,
                   color=None, bold=None):
    style.font.name = ascii_name
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_shading(paragraph, fill, color=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if color:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(color)


def set_cell_border(cell, color="C9D2DC", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, size=11, color=TEXT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.widow_control = True

    h1 = styles["Heading 1"]
    set_style_font(h1, size=16, color=BLUE, bold=True)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.page_break_before = True

    h2 = styles["Heading 2"]
    set_style_font(h2, size=13, color=BLUE, bold=True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    set_style_font(h3, size=12, color=DARK_BLUE, bold=True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        set_style_font(style, size=11, color=TEXT)
        pf = style.paragraph_format
        pf.left_indent = Inches(0.375)
        pf.first_line_indent = Inches(-0.194)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.208

    caption = styles["Caption"]
    set_style_font(caption, size=9, color=MUTED)
    caption.font.italic = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_together = True

    if "Report Kicker" not in styles:
        s = styles.add_style("Report Kicker", WD_STYLE_TYPE.PARAGRAPH)
    else:
        s = styles["Report Kicker"]
    set_style_font(s, size=10, color=GOLD, bold=True)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(18)

    if "Report Title" not in styles:
        s = styles.add_style("Report Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        s = styles["Report Title"]
    set_style_font(s, size=27, color=NAVY, bold=True)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(14)
    s.paragraph_format.keep_together = True

    if "Report Subtitle" not in styles:
        s = styles.add_style("Report Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        s = styles["Report Subtitle"]
    set_style_font(s, size=15, color=BLUE)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(20)
    s.paragraph_format.keep_together = True

    if "Report Meta" not in styles:
        s = styles.add_style("Report Meta", WD_STYLE_TYPE.PARAGRAPH)
    else:
        s = styles["Report Meta"]
    set_style_font(s, size=10, color=MUTED)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(6)

    if "Callout" not in styles:
        s = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        s = styles["Callout"]
    set_style_font(s, size=10.5, color=DARK_BLUE)
    s.paragraph_format.left_indent = Inches(0.16)
    s.paragraph_format.right_indent = Inches(0.16)
    s.paragraph_format.space_before = Pt(6)
    s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.line_spacing = 1.2


def configure_section(section, *, cover=False):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    sect_pr = section._sectPr
    v_align = sect_pr.find(qn("w:vAlign"))
    if v_align is None:
        v_align = OxmlElement("w:vAlign")
        sect_pr.append(v_align)
    v_align.set(qn("w:val"), "center" if cover else "top")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9, color=MUTED)


def configure_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("MS-CADM 复现与改进实验综合报告")
    set_run_font(r, size=9, color=MUTED, bold=True)
    r2 = p.add_run("    ·    2026-08-09")
    set_run_font(r2, size=8.5, color="8A949C")

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("multi_scaleCADM 研究档案")
    set_run_font(r, size=8.5, color="8A949C")
    p.add_run("                                                          ")
    add_page_number(p)

    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), "1")


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\(https?://[^\s)]+\)|https?://[^\s]+)")


def add_inline(paragraph, text, *, size=None, color=None):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, ascii_name="Consolas", east_asia="Microsoft YaHei", size=(size or 10), color=DARK_BLUE)
            r_pr = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF2F5")
            r_pr.append(shd)
        elif token.startswith("["):
            link_match = re.fullmatch(r"\[([^\]]+)\]\((https?://[^\s)]+)\)", token)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
            else:
                run = paragraph.add_run(token)
                set_run_font(run, size=size, color=color)
        else:
            trailing = ""
            while token and token[-1] in ".,;:，。；：)）":
                trailing = token[-1] + trailing
                token = token[:-1]
            add_hyperlink(paragraph, token, token)
            if trailing:
                run = paragraph.add_run(trailing)
                set_run_font(run, size=size, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=color)


def add_cover(doc: Document, cover_lines: list[str]):
    cover_section = doc.sections[0]
    configure_section(cover_section, cover=True)
    cover_section.header.is_linked_to_previous = False
    cover_section.footer.is_linked_to_previous = False
    cover_section.header.paragraphs[0].clear()
    cover_section.footer.paragraphs[0].clear()

    p = doc.add_paragraph(style="Report Kicker")
    p.add_run("RESEARCH REPORT  ·  EVIDENCE AUDIT  ·  2026")

    title = cover_lines[0].lstrip("# ").strip()
    p = doc.add_paragraph(style="Report Title")
    if title.endswith("综合研究报告"):
        first = title[:-len("综合研究报告")].rstrip()
        p.add_run(first)
        p.add_run().add_break()
        p.add_run("综合研究报告")
    else:
        p.add_run(title)

    subtitle = cover_lines[1].strip()
    p = doc.add_paragraph(style="Report Subtitle")
    p.add_run(subtitle)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    set_run_font(r, size=9, color=GOLD)

    for line in cover_lines[2:]:
        if not line.strip():
            continue
        p = doc.add_paragraph(style="Report Meta")
        p.add_run(line.strip())

    p = doc.add_paragraph(style="Report Meta")
    p.paragraph_format.space_before = Pt(26)
    r = p.add_run("包含：正式复现 · 探索实验 · 内部确认 · exact SUC · 负结果与审计")
    set_run_font(r, size=9.5, color=DARK_BLUE, bold=True)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section, cover=False)
    configure_header_footer(body_section)


def font_path(bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf" if bold else r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return str(path)
    return None


def get_font(size, bold=False):
    path = font_path(bold)
    return ImageFont.truetype(path, size=size) if path else ImageFont.load_default()


def draw_wrapped(draw, xy, text, font, fill, max_width, line_gap=8, anchor="la"):
    x, y = xy
    lines = []
    current = ""
    for ch in text:
        trial = current + ch
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = ch
    if current:
        lines.append(current)
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill, anchor=anchor)
        y += font.size + line_gap
    return y


def create_timeline(path: Path):
    w, h = 1800, 1000
    img = Image.new("RGB", (w, h), "#" + WHITE)
    d = ImageDraw.Draw(img)
    title_f = get_font(46, True)
    sub_f = get_font(25)
    node_f = get_font(29, True)
    small_f = get_font(21)
    d.text((90, 72), "从复现到安全决策：七阶段研究链", font=title_f, fill="#" + NAVY)
    d.text((90, 132), "颜色表示当前证据状态；箭头表示问题与方法的递进，而非统一 leaderboard。", font=sub_f, fill="#" + MUTED)
    y_line = 505
    d.line((130, y_line, 1670, y_line), fill="#B7C2CC", width=8)
    nodes = [
        (150, "Baseline", "正式复现\nheadline 未重现", BLUE_GREY),
        (400, "CR", "探索性强改善\n修复欠离散", AMBER),
        (650, "RAHC", "4/11 门槛\n尾部过扩张", RED),
        (900, "CAA", "9/9 同向\ncoverage 未达标", AMBER),
        (1150, "MM", "当前最强正结果\n混合测度有效", GREEN),
        (1400, "STGF", "6/6 门槛失败\n低 NFE 有竞争力", RED),
        (1650, "PS", "安全回退成功\n决策收益为 0", RED),
    ]
    for i, (x, name, desc, color) in enumerate(nodes):
        d.ellipse((x - 27, y_line - 27, x + 27, y_line + 27), fill="#" + color, outline="#" + WHITE, width=5)
        card_w, card_h = 218, 238
        top = 225 if i % 2 == 0 else 570
        left = x - card_w // 2
        d.rounded_rectangle((left, top, left + card_w, top + card_h), radius=22,
                            fill="#F7F9FB", outline="#" + color, width=5)
        d.text((x, top + 42), name, font=node_f, fill="#" + color, anchor="ma")
        yy = top + 98
        for line in desc.split("\n"):
            d.text((x, yy), line, font=small_f, fill="#" + TEXT, anchor="ma")
            yy += 39
        endpoint_y = top + card_h if top < y_line else top
        d.line((x, endpoint_y, x, y_line - 30 if top < y_line else y_line + 30), fill="#" + color, width=4)
    d.text((90, 920), "证据等级：内部确认 ≠ 外部验证；CR/RAHC 与小样本 SUC 均不得用于确认性主张。",
           font=get_font(24, True), fill="#" + DARK_BLUE)
    img.save(path, quality=95)


def create_ps_gate_matrix(path: Path):
    w, h = 1800, 900
    img = Image.new("RGB", (w, h), "#" + WHITE)
    d = ImageDraw.Draw(img)
    d.text((85, 55), "PS-DFSC：12 个验证候选均未通过安全发布门槛", font=get_font(44, True), fill="#" + NAVY)
    d.text((85, 115), "主要共同失败为 90% coverage < 0.88；红色叉号表示不能进入 exact validation shortlist。",
           font=get_font(24), fill="#" + MUTED)
    coverages = [
        [0.874667, 0.866083, 0.865417, 0.859917],
        [0.871333, 0.862333, 0.861000, 0.861333],
        [0.866750, 0.867333, 0.866917, 0.865000],
    ]
    fallbacks = [[9, 20, 19, 32], [17, 19, 26, 30], [24, 25, 35, 36]]
    betas = ["β=0", "β=0.25", "β=0.5", "β=1"]
    left, top = 260, 225
    cw, ch = 345, 170
    for j, beta in enumerate(betas):
        d.text((left + j * cw + cw / 2, top - 50), beta, font=get_font(27, True), fill="#" + DARK_BLUE, anchor="ma")
    for i in range(3):
        d.text((175, top + i * ch + ch / 2), f"Outer {i + 1}", font=get_font(27, True), fill="#" + DARK_BLUE, anchor="mm")
        for j in range(4):
            x0 = left + j * cw
            y0 = top + i * ch
            d.rounded_rectangle((x0 + 10, y0 + 10, x0 + cw - 10, y0 + ch - 10), radius=18,
                                fill="#FCEDEF", outline="#" + RED, width=4)
            d.text((x0 + 42, y0 + 52), "×", font=get_font(54, True), fill="#" + RED, anchor="mm")
            d.text((x0 + 85, y0 + 37), f"coverage  {coverages[i][j]:.4f}", font=get_font(24, True), fill="#" + TEXT)
            d.text((x0 + 85, y0 + 82), f"daily fallback  {fallbacks[i][j]}/50", font=get_font(21), fill="#" + MUTED)
            d.text((x0 + 85, y0 + 119), "candidate = unsafe", font=get_font(20), fill="#" + RED)
    d.rounded_rectangle((260, 770, 1640, 845), radius=16, fill="#EDF4FA", outline="#" + BLUE, width=3)
    d.text((950, 808), "结果：3 个 outer 均锁定 identity → exact 确认成本差 = 0；fail-closed 机制按设计工作。",
           font=get_font(24, True), fill="#" + NAVY, anchor="mm")
    img.save(path, quality=95)


def create_custom_figure(name: str) -> Path:
    if name == "timeline":
        path = ASSET_DIR / "research_timeline.png"
        create_timeline(path)
        return path
    if name == "ps_gate_matrix":
        path = ASSET_DIR / "ps_dfsc_gate_matrix.png"
        create_ps_gate_matrix(path)
        return path
    raise ValueError(f"Unknown custom figure: {name}")


def add_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_figure(doc: Document, path: Path, caption: str):
    if not path.exists():
        p = doc.add_paragraph(style="Callout")
        add_inline(p, f"图像缺失：{path}")
        set_paragraph_shading(p, "FCEDEF")
        return
    with Image.open(path) as im:
        px_w, px_h = im.size
    max_w, max_h = (5.5, 5.8) if path.name == "research_timeline.png" else (6.3, 5.8)
    aspect = px_w / px_h
    width = max_w
    height = width / aspect
    if height > max_h:
        height = max_h
        width = height * aspect
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width), height=Inches(height))
    add_alt_text(shape, caption.split("。")[0], caption)
    cp = doc.add_paragraph(style="Caption")
    add_inline(cp, caption, size=9, color=MUTED)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def table_weights(rows: list[list[str]]) -> list[float]:
    n = len(rows[0])
    if n == 2:
        return [2.1, 4.4]
    if n == 3:
        return [2.0, 1.35, 3.15]
    if n == 4:
        return [1.65, 1.15, 1.15, 2.55]
    if n == 5:
        return [1.55, 1.1, 1.1, 1.1, 1.65]
    if n == 6:
        return [1.4, 1.0, 1.0, 1.0, 1.0, 1.1]
    if n == 7:
        return [1.25, 0.85, 0.85, 0.85, 0.85, 0.85, 1.0]
    if n == 8:
        return [1.0] + [0.78] * 7
    return [1.0] * n


NUMERIC_RE = re.compile(r"^[+−-]?[\d,.]+(?:%|秒)?(?:\s*[/→±].*)?$|^—$|^\d+/\d+$")


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    ncols = len(rows[0])
    if any(len(row) != ncols for row in rows):
        raise ValueError(f"Ragged markdown table: {rows[:2]}")
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = 0
    for r_idx, row in enumerate(rows):
        tr = table.rows[r_idx]
        set_row_cant_split(tr)
        for c_idx, value in enumerate(row):
            cell = tr.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else (
                WD_ALIGN_PARAGRAPH.RIGHT if NUMERIC_RE.fullmatch(value.strip()) else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline(p, value, size=8.3 if ncols >= 6 else 8.8, color=WHITE if r_idx == 0 else TEXT)
            for run in p.runs:
                if r_idx == 0:
                    run.bold = True
            set_cell_border(cell)
            if r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F7F9FB")
        if r_idx == 0:
            set_repeat_table_header(tr)
    widths = column_widths_from_weights(table_weights(rows), total_width_dxa=9360)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=9360,
        indent_dxa=120,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120},
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)


def parse_body(doc: Document, lines: list[str]):
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line == "[[END_REPORT]]":
            break
        if line.startswith("[[FIGURE:"):
            payload = line[len("[[FIGURE:"):-2]
            rel, caption = payload.split("|", 1)
            add_figure(doc, ROOT / rel, caption)
            i += 1
            continue
        if line.startswith("[[CUSTOM_FIGURE:"):
            payload = line[len("[[CUSTOM_FIGURE:"):-2]
            name, caption = payload.split("|", 1)
            add_figure(doc, create_custom_figure(name), caption)
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[2:].strip())
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[3:].strip())
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, line[4:].strip())
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and is_separator_row(lines[i + 1]):
            rows = [split_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:].strip())
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph(style="Callout")
            add_inline(p, line[2:].strip())
            set_paragraph_shading(p, LIGHT_BLUE)
            i += 1
            continue
        p = doc.add_paragraph(style="Normal")
        add_inline(p, line)
        i += 1


def set_document_compatibility(doc: Document):
    settings = doc.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    compat_setting = OxmlElement("w:compatSetting")
    compat_setting.set(qn("w:name"), "compatibilityMode")
    compat_setting.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    compat_setting.set(qn("w:val"), "15")
    compat.append(compat_setting)


def main():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    marker = lines.index("[[COVER_END]]")
    cover_lines = [line for line in lines[:marker] if line.strip()]
    body_lines = lines[marker + 1:]

    doc = Document()
    configure_styles(doc)
    set_document_compatibility(doc)
    props = doc.core_properties
    props.title = "MS-CADM 复现与改进实验综合研究报告"
    props.subject = "GEFCom2014 风电概率场景生成与随机机组组合"
    props.author = "multi_scaleCADM research workspace"
    props.keywords = "MS-CADM, wind scenarios, calibration, MM-JDWind, STGF-Flow, PS-DFSC, SUC"
    props.comments = "Generated from an audited Markdown source; report date 2026-08-09."

    add_cover(doc, cover_lines)
    parse_body(doc, body_lines)

    for section in doc.sections:
        if section.start_type != WD_SECTION_START.NEW_PAGE and section is not doc.sections[0]:
            section.start_type = WD_SECTION_START.NEW_PAGE

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Inline shapes: {len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
