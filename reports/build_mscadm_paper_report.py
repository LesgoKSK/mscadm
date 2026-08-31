from __future__ import annotations

import json
import math
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_comprehensive_report as base


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "MSCADM_PAPER_REPRODUCTION_REPORT.md"
OUTPUT = ROOT / "reports" / "MS-CADM_论文方法与复现专题报告.docx"
ASSET_DIR = ROOT / "reports" / "assets" / "mscadm_paper"
ASSET_DIR.mkdir(parents=True, exist_ok=True)


def configure_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("MS-CADM 论文方法与复现专题报告")
    base.set_run_font(run, size=9, color=base.MUTED, bold=True)
    run = paragraph.add_run("    ·    2026-08-09")
    base.set_run_font(run, size=8.5, color="8A949C")

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run("论文原理 · 代码映射 · 复现证据")
    base.set_run_font(run, size=8.5, color="8A949C")
    paragraph.add_run("                                                          ")
    base.add_page_number(paragraph)

    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), "1")


def add_cover(doc: Document, cover_lines: list[str]) -> None:
    cover_section = doc.sections[0]
    base.configure_section(cover_section, cover=True)
    cover_section.header.is_linked_to_previous = False
    cover_section.footer.is_linked_to_previous = False
    cover_section.header.paragraphs[0].clear()
    cover_section.footer.paragraphs[0].clear()

    paragraph = doc.add_paragraph(style="Report Kicker")
    paragraph.add_run("PAPER REPRODUCTION  ·  ARCHITECTURE AUDIT  ·  2026")

    title = cover_lines[0].lstrip("# ").strip()
    paragraph = doc.add_paragraph(style="Report Title")
    if "：" in title:
        first, second = title.split("：", 1)
        paragraph.add_run(first + "：")
        paragraph.add_run().add_break()
        paragraph.add_run(second)
    else:
        paragraph.add_run(title)

    paragraph = doc.add_paragraph(style="Report Subtitle")
    paragraph.add_run(cover_lines[1].strip())

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    base.set_run_font(run, size=9, color=base.GOLD)

    for line in cover_lines[2:]:
        if not line.strip():
            continue
        paragraph = doc.add_paragraph(style="Report Meta")
        paragraph.add_run(line.strip())

    paragraph = doc.add_paragraph(style="Report Meta")
    paragraph.paragraph_format.space_before = Pt(24)
    run = paragraph.add_run("覆盖：论文原理 · 网络张量流 · 训练/采样 · 数据协议 · 主表/消融/SUC · 复现边界")
    base.set_run_font(run, size=9.4, color=base.DARK_BLUE, bold=True)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    base.configure_section(body_section, cover=False)
    configure_header_footer(body_section)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int],
          *, color: str = "#607D8B", width: int = 8) -> None:
    draw.line((start, end), fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 22
    spread = 0.55
    p1 = (
        end[0] - length * math.cos(angle - spread),
        end[1] - length * math.sin(angle - spread),
    )
    p2 = (
        end[0] - length * math.cos(angle + spread),
        end[1] - length * math.sin(angle + spread),
    )
    draw.polygon((end, p1, p2), fill=color)


def rounded_box(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], title: str,
                lines: list[str], *, fill: str, outline: str, title_color: str | None = None,
                title_size: int = 38, text_size: int = 28) -> None:
    x0, y0, x1, y1 = xy
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=outline, width=5)
    draw.text(((x0 + x1) // 2, y0 + 42), title, font=base.get_font(title_size, True),
              fill=title_color or outline, anchor="mm")
    y = y0 + 88
    font = base.get_font(text_size)
    for line in lines:
        draw.text(((x0 + x1) // 2, y), line, font=font, fill="#263238", anchor="ma")
        y += text_size + 15


def create_architecture(path: Path) -> None:
    width, height = 2200, 1450
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((90, 58), "MS-CADM 复现实现：端到端张量流", font=base.get_font(52, True), fill="#17365D")
    draw.text((90, 122), "实线为前向数据流；蓝色框为条件路径，金色框为扩散时间步调制。",
              font=base.get_font(27), fill="#66737F")

    rounded_box(draw, (70, 250, 410, 440), "NWP + Zone", ["[B, 24, 20]", "10 NWP + 10 one-hot"],
                fill="#EAF2F8", outline="#2E74B5")
    rounded_box(draw, (70, 565, 410, 735), "Noisy target x_t", ["[B, 24, 1]", "每小时一个 token"],
                fill="#F4F6F9", outline="#607D8B")
    rounded_box(draw, (70, 875, 410, 1035), "Diffusion step t", ["[B]", "0 … 249"],
                fill="#FFF6E5", outline="#C59A3D", text_size=24)

    rounded_box(draw, (520, 205, 995, 490), "Multi-scale condition encoder",
                ["1×1: 20 → 128", "down: 128 → 64", "(3, 5, 7) 双卷积分支", "concat 192 → 64 → 128", "residual + LayerNorm"],
                fill="#EAF2F8", outline="#2E74B5", title_size=30, text_size=24)
    rounded_box(draw, (535, 565, 980, 755), "Target token embedding",
                ["Linear 1 → 128", "+ learned position [1,24,128]"],
                fill="#F4F6F9", outline="#607D8B", title_size=31)
    rounded_box(draw, (535, 860, 980, 1060), "Timestep embedding",
                ["sin/cos 128", "MLP 128 → 512 → 128"],
                fill="#FFF6E5", outline="#C59A3D", title_size=35)

    arrow(draw, (410, 345), (520, 345), color="#2E74B5")
    arrow(draw, (410, 650), (535, 650), color="#607D8B")
    arrow(draw, (410, 950), (535, 950), color="#C59A3D")

    rounded_box(draw, (1100, 275, 1780, 785), "4 × AdaLN Transformer block",
                ["hidden [B,24,128]", "LN(no affine) → scale/shift", "4-head self-attention (head=32)", "gate 1 + residual", "LN → scale/shift", "FFN 128 → 512 → 128", "gate 2 + residual", "time MLP 输出 6×128 调制量"],
                fill="#F7F9FB", outline="#1F4D78", title_size=39, text_size=29)

    draw.ellipse((1015, 490, 1085, 560), fill="#FFFFFF", outline="#2E74B5", width=5)
    draw.text((1050, 525), "+", font=base.get_font(42, True), fill="#2E74B5", anchor="mm")
    arrow(draw, (995, 345), (1040, 490), color="#2E74B5")
    arrow(draw, (980, 650), (1040, 560), color="#607D8B")
    arrow(draw, (1085, 525), (1100, 525), color="#1F4D78")
    arrow(draw, (980, 950), (1320, 785), color="#C59A3D")

    rounded_box(draw, (1860, 330, 2140, 505), "Noise head", ["eps [B,24,1]"],
                fill="#EDF7EF", outline="#2E7D32", title_size=34, text_size=28)
    rounded_box(draw, (1860, 575, 2140, 750), "Variance head", ["raw_var [B,24,1]"],
                fill="#FCEDEF", outline="#B23A48", title_size=34, text_size=28)
    arrow(draw, (1780, 465), (1860, 420), color="#2E7D32")
    arrow(draw, (1780, 625), (1860, 660), color="#B23A48")

    draw.rounded_rectangle((80, 1175, 2120, 1360), radius=22, fill="#EEF4FA", outline="#2E74B5", width=4)
    draw.text((120, 1215), "复现锁定值", font=base.get_font(31, True), fill="#17365D")
    draw.text((360, 1215), "dim=128 · bottleneck=64 · depth=4 · heads=4 · FF×4 · 1,478,018 parameters",
              font=base.get_font(29, True), fill="#263238")
    draw.text((120, 1270), "论文明确：3/5/7 卷积、AdaLN、双输出；论文未披露：上述维度、层数、head、初始化和 patch 细节。",
              font=base.get_font(27), fill="#1F4D78")
    image.save(path, quality=95)


def create_diffusion_workflow(path: Path) -> None:
    width, height = 2000, 1020
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((85, 55), "训练与生成：本地采用自洽的 Improved-DDPM 解释",
              font=base.get_font(48, True), fill="#17365D")
    draw.text((85, 117), "论文 Algorithm 2 的逐步重新采样无法形成反向链；复现实现改为 x_t → x_{t-1} 的递推。",
              font=base.get_font(27), fill="#66737F")

    draw.text((90, 220), "TRAIN", font=base.get_font(34, True), fill="#2E74B5")
    boxes = [
        ((190, 270, 480, 450), "真实轨迹 x_0", ["标准化 [B,24,1]"]),
        ((650, 270, 995, 450), "随机 t 与 eps", ["t ~ Uniform(0,T-1)", "eps ~ N(0,I)"]),
        ((1170, 270, 1515, 450), "构造 x_t", ["sqrt(alpha_bar_t) x_0 +", "sqrt(1-alpha_bar_t) eps"]),
        ((1680, 270, 1940, 450), "MS-CADM", ["预测 eps 与 raw_var"]),
    ]
    for xy, title, lines in boxes:
        rounded_box(draw, xy, title, lines, fill="#EAF2F8", outline="#2E74B5", title_size=31, text_size=24)
    for a, b in zip(boxes[:-1], boxes[1:]):
        arrow(draw, (a[0][2], 360), (b[0][0], 360), color="#2E74B5")
    draw.rounded_rectangle((620, 495, 1770, 575), radius=18, fill="#F4F6F9", outline="#607D8B", width=3)
    draw.text((1195, 535), "L = MSE(eps_pred, eps) + 1e-3 · VLB_var（均值路径 stop-gradient）",
              font=base.get_font(27, True), fill="#263238", anchor="mm")

    draw.text((90, 650), "SAMPLE", font=base.get_font(34, True), fill="#2E7D32")
    samples = [
        ((190, 700, 500, 875), "一次采样 x_T", ["x_T ~ N(0,I)"]),
        ((710, 700, 1120, 875), "递推 249 → 0", ["预测 eps, log_var", "x_t → x_{t-1}"]),
        ((1330, 700, 1640, 875), "逆标准化", ["clip 到 [0,1]"]),
        ((1770, 700, 1950, 875), "100 条", ["场景/日"]),
    ]
    for xy, title, lines in samples:
        rounded_box(draw, xy, title, lines, fill="#EDF7EF", outline="#2E7D32", title_size=30, text_size=24)
    for a, b in zip(samples[:-1], samples[1:]):
        arrow(draw, (a[0][2], 790), (b[0][0], 790), color="#2E7D32")
    image.save(path, quality=95)


def create_reproduction_scorecard(path: Path) -> None:
    width, height = 1900, 980
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((80, 55), "MS-CADM headline：本地值 / 论文值", font=base.get_font(48, True), fill="#17365D")
    draw.text((80, 115), "六项指标均越低越好；100% 表示数值完全一致，超过 100% 表示本地复现更差。",
              font=base.get_font(27), fill="#66737F")
    labels = ["MAE", "RMSE", "CRPS", "QS", "ES", "VS"]
    ratios = [104.32, 107.98, 123.02, 122.76, 126.76, 121.39]
    paper = [0.1191, 0.1645, 0.0873, 0.0441, 0.5380, 18.14]
    local = [0.124248, 0.177632, 0.107400, 0.054139, 0.681950, 22.0200]
    left, top, right, bottom = 620, 235, 1760, 830
    draw.line((left, bottom, right, bottom), fill="#9BA8B4", width=4)
    x100 = left + (100 - 95) / (130 - 95) * (right - left)
    draw.line((x100, top, x100, bottom), fill="#2E7D32", width=4)
    draw.text((x100, top - 28), "100%", font=base.get_font(24, True), fill="#2E7D32", anchor="ms")
    for i, (label, ratio, p, r) in enumerate(zip(labels, ratios, paper, local)):
        y = top + 55 + i * 92
        draw.text((80, y), label, font=base.get_font(31, True), fill="#1F4D78", anchor="lm")
        draw.text((185, y), f"{p:.4f} → {r:.4f}", font=base.get_font(25), fill="#66737F", anchor="lm")
        x = left + (ratio - 95) / (130 - 95) * (right - left)
        color = "#C27C0E" if ratio < 115 else "#B23A48"
        draw.line((left, y, x, y), fill=color, width=34)
        draw.ellipse((x - 18, y - 18, x + 18, y + 18), fill=color)
        draw.text((x + 32, y), f"{ratio:.2f}%", font=base.get_font(27, True), fill=color, anchor="lm")
    draw.text((80, 915), "最主要偏差集中在概率分数：CRPS、QS、ES、VS 相对退化约 21%–27%。",
              font=base.get_font(29, True), fill="#B23A48")
    image.save(path, quality=95)


def create_training_curve(path: Path) -> None:
    history_path = ROOT / "outputs" / "full_reproduction" / "mscadm" / "history.json"
    history = json.loads(history_path.read_text(encoding="utf-8"))
    steps = [float(row["step"]) for row in history]
    losses = [float(row["loss"]) for row in history]
    simple = [float(row["simple"]) for row in history]
    width, height = 1900, 930
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((80, 55), "MS-CADM 26,000-step 训练轨迹", font=base.get_font(48, True), fill="#17365D")
    draw.text((80, 115), "记录间隔 100 step；总训练墙钟约 2,293 秒（RTX 2060）。",
              font=base.get_font(27), fill="#66737F")
    left, top, right, bottom = 170, 220, 1780, 780
    draw.rectangle((left, top, right, bottom), outline="#B7C2CC", width=3)
    ymin, ymax = min(min(losses), min(simple)), max(max(losses), max(simple))
    ymin = math.floor(ymin * 20) / 20
    ymax = math.ceil(ymax * 20) / 20
    for j in range(6):
        value = ymin + (ymax - ymin) * j / 5
        y = bottom - (value - ymin) / (ymax - ymin) * (bottom - top)
        draw.line((left, y, right, y), fill="#E1E6EA", width=2)
        draw.text((left - 18, y), f"{value:.2f}", font=base.get_font(22), fill="#66737F", anchor="rm")
    for j in range(6):
        value = 26000 * j / 5
        x = left + value / 26000 * (right - left)
        draw.text((x, bottom + 22), f"{int(value/1000)}k", font=base.get_font(22), fill="#66737F", anchor="ma")

    def points(values: list[float]) -> list[tuple[float, float]]:
        return [
            (
                left + step / 26000 * (right - left),
                bottom - (value - ymin) / (ymax - ymin) * (bottom - top),
            )
            for step, value in zip(steps, values)
        ]

    draw.line(points(losses), fill="#2E74B5", width=5)
    draw.line(points(simple), fill="#C59A3D", width=3)
    draw.line((1320, 160, 1410, 160), fill="#2E74B5", width=7)
    draw.text((1425, 160), "total", font=base.get_font(24, True), fill="#2E74B5", anchor="lm")
    draw.line((1550, 160, 1640, 160), fill="#C59A3D", width=5)
    draw.text((1655, 160), "simple", font=base.get_font(24, True), fill="#C59A3D", anchor="lm")
    draw.text((80, 865), f"首条记录 loss={losses[0]:.4f}；末条记录 loss={losses[-1]:.4f}。下降并不等价于概率校准良好。",
              font=base.get_font(27, True), fill="#1F4D78")
    image.save(path, quality=95)


def create_custom_figure(name: str) -> Path:
    builders = {
        "architecture": (ASSET_DIR / "mscadm_architecture.png", create_architecture),
        "diffusion_workflow": (ASSET_DIR / "diffusion_workflow.png", create_diffusion_workflow),
        "reproduction_scorecard": (ASSET_DIR / "reproduction_scorecard.png", create_reproduction_scorecard),
        "training_curve": (ASSET_DIR / "training_curve.png", create_training_curve),
    }
    if name not in builders:
        raise ValueError(f"Unknown custom figure: {name}")
    path, builder = builders[name]
    builder(path)
    return path


def main() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    marker = lines.index("[[COVER_END]]")
    cover_lines = [line for line in lines[:marker] if line.strip()]
    body_lines = lines[marker + 1 :]

    base.configure_header_footer = configure_header_footer
    base.create_custom_figure = create_custom_figure

    doc = Document()
    base.configure_styles(doc)
    base.set_document_compatibility(doc)
    props = doc.core_properties
    props.title = "MS-CADM 论文方法与复现专题报告"
    props.subject = "MS-CADM 网络架构、扩散原理、GEFCom2014 复现与证据审计"
    props.author = "multi_scaleCADM research workspace"
    props.keywords = "MS-CADM, diffusion transformer, wind scenarios, GEFCom2014, reproduction"
    props.comments = "Generated from an audited Markdown source; report date 2026-08-09."

    add_cover(doc, cover_lines)
    base.parse_body(doc, body_lines)

    # Keep the compact SUC reconstruction subsection together on its own page.
    # Without this explicit break, the final explanatory paragraph can be left
    # alone immediately before the next chapter's mandatory page break.
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("8.7 Table 5：RTS-24 SUC 重构"):
            paragraph.paragraph_format.page_break_before = True
            break

    for section in doc.sections:
        if section.start_type != WD_SECTION_START.NEW_PAGE and section is not doc.sections[0]:
            section.start_type = WD_SECTION_START.NEW_PAGE

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Inline shapes: {len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
